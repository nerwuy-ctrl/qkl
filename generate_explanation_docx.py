from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "项目功能实现与使用说明.md"
TARGET = ROOT / "供应链金融平台项目功能实现与使用说明.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])


def add_inline(paragraph, text):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
            run.font.color.rgb = RGBColor(30, 105, 78)
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.3)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "宋体"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.45
normal.paragraph_format.space_after = Pt(5)

for name, size, color in [
    ("Title", 24, "173D31"),
    ("Heading 1", 18, "176B4D"),
    ("Heading 2", 15, "176B4D"),
    ("Heading 3", 12, "264E40"),
]:
    style = styles[name]
    style.font.name = "微软雅黑"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True

header = section.header.paragraphs[0]
header.text = "供应链金融平台"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(9)
header.runs[0].font.color.rgb = RGBColor(100, 120, 110)
add_page_number(section.footer.paragraphs[0])

lines = SOURCE.read_text(encoding="utf-8").splitlines()
in_code = False
code_lines = []
i = 0

while i < len(lines):
    line = lines[i]
    stripped = line.strip()

    if stripped.startswith("```"):
        if not in_code:
            in_code = True
            code_lines = []
        else:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.right_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(7)
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
            run.font.size = Pt(9)
            set_cell_shading_dummy = OxmlElement("w:shd")
            set_cell_shading_dummy.set(qn("w:fill"), "F1F5F2")
            p._p.get_or_add_pPr().append(set_cell_shading_dummy)
            in_code = False
        i += 1
        continue

    if in_code:
        code_lines.append(line)
        i += 1
        continue

    if stripped == "---":
        i += 1
        continue

    if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[i + 1]):
        table_lines = [line]
        i += 2
        while i < len(lines) and lines[i].strip().startswith("|"):
            table_lines.append(lines[i])
            i += 1
        rows = [[c.strip() for c in x.strip().strip("|").split("|")] for x in table_lines]
        cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        for r_index, row_data in enumerate(rows):
            row = table.rows[r_index]
            if r_index == 0:
                set_repeat_table_header(row)
            for c_index, value in enumerate(row_data):
                cell = row.cells[c_index]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = ""
                add_inline(cell.paragraphs[0], value)
                if r_index == 0:
                    set_cell_shading(cell, "DCEDE4")
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(9)
        doc.add_paragraph().paragraph_format.space_after = Pt(1)
        continue

    heading = re.match(r"^(#{1,6})\s+(.*)", stripped)
    if heading:
        level = len(heading.group(1))
        text = heading.group(2)
        if level == 1:
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(10)
            p.add_run(text)
        else:
            doc.add_heading(text, level=min(level - 1, 3))
        i += 1
        continue

    numbered = re.match(r"^(\d+)\.\s+(.*)", stripped)
    bullet = re.match(r"^-\s+(.*)", stripped)
    if numbered:
        p = doc.add_paragraph(style="List Number")
        add_inline(p, numbered.group(2))
    elif bullet:
        p = doc.add_paragraph(style="List Bullet")
        add_inline(p, bullet.group(1))
    elif stripped:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        add_inline(p, stripped)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    i += 1

doc.core_properties.title = "供应链金融平台项目功能实现与使用说明"
doc.core_properties.subject = "区块链供应链金融平台课程项目说明"
doc.core_properties.author = "王宇诺"
doc.save(TARGET)
print(TARGET)
